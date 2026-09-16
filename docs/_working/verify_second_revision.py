from pathlib import Path
from docx import Document

SOURCE = Path(r"C:\xampp\htdocs\smartqms\docs\Grp25-IT225-Chapters123-widcomments-July9-2026-updated.docx")
document = Document(SOURCE)

def index_of(text):
    return next(index for index, paragraph in enumerate(document.paragraphs) if paragraph.text.strip() == text)

for caption in (
    "Figure 3: Context Diagram",
    "Figure 4: Data Flow Diagram",
    "Figure 5: Entity Relationship Diagram",
    "Figure 6: Use Case Diagram",
):
    previous = document.paragraphs[index_of(caption) - 1]
    print(caption, "picture_before_caption=", len(previous._p.xpath(".//w:drawing")) == 1)

entity_start = index_of("Entity Descriptions:") + 1
use_case = index_of("Use Case Diagram")
entities = [paragraph.text.strip() for paragraph in document.paragraphs[entity_start:use_case]]
print("entity_descriptions", len(entities), [text.split(" - ")[0] for text in entities])

development = index_of("Development Tools")
client_start = index_of("Client Use Cases:")
use_cases = [paragraph.text.strip() for paragraph in document.paragraphs[client_start:development]]
print("use_case_headings", [text for text in use_cases if text.endswith("Use Cases:")])
print("use_case_count", len(use_cases))

tool_start = index_of("Evaluation Tool")
procedure_start = index_of("Evaluation Procedure")
procedure = [paragraph.text.strip() for paragraph in document.paragraphs[procedure_start + 1:tool_start]]
print("procedure_count", len(procedure), "contains_maintainability", any("Maintainability" in text for text in procedure))

table = next(table for table in document.tables if table.cell(0, 0).text.strip() == "CRITERIA")
criteria = [table.cell(row, 0).text.strip() for row in range(len(table.rows))]
print("criteria", criteria)
print("has_maintainability", any("Maintainability" in value for value in criteria))
print("has_placeholder", any("what item" in value or "define according" in value for value in criteria + procedure))
print("inline_shapes", len(document.inline_shapes))
