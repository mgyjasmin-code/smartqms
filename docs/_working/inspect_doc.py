from docx import Document

SOURCE = r"C:\xampp\htdocs\smartqms\docs\_working\Grp25-IT225-Chapters123-widcomments-July9-2026.original.docx"
NEEDLES = (
    "Data Flow Diagram",
    "paste the",
    "below is an example",
    "Evaluation Procedure",
    "Functional Sustainability",
    "Evaluation Tool",
    "what item in",
)

document = Document(SOURCE)
for index, paragraph in enumerate(document.paragraphs):
    text = paragraph.text.strip()
    if any(needle in text for needle in NEEDLES):
        drawings = len(paragraph._p.xpath(".//w:drawing"))
        print(index, repr(text), f"drawings={drawings}")

for table_index, table in enumerate(document.tables):
    for row_index, row in enumerate(table.rows):
        for cell_index, cell in enumerate(row.cells):
            for paragraph_index, paragraph in enumerate(cell.paragraphs):
                text = paragraph.text.strip()
                if any(needle in text for needle in NEEDLES):
                    print(
                        f"table={table_index} row={row_index} cell={cell_index} paragraph={paragraph_index}",
                        repr(text),
                    )

print("inline_shapes", len(document.inline_shapes))
for index, paragraph in enumerate(document.paragraphs):
    drawings = len(paragraph._p.xpath(".//w:drawing"))
    if drawings:
        print("drawing paragraph", index, repr(paragraph.text.strip()), drawings)

for index in range(198, 215):
    print("context", index, repr(document.paragraphs[index].text.strip()), "drawings=", len(document.paragraphs[index]._p.xpath(".//w:drawing")))

for index in range(248, 257):
    print("evaluation", index, repr(document.paragraphs[index].text.strip()))
