from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, RGBColor

SOURCE = Path(r"C:\xampp\htdocs\smartqms\docs\_working\Grp25-rebased.docx")
OUTPUT = Path(r"C:\xampp\htdocs\smartqms\docs\Grp25-IT225-Chapters123-widcomments-July9-2026-updated.docx")
DIAGRAMS = Path(r"C:\xampp\htdocs\smartqms\docs\diagrams")

FIGURES = {
    "Figure 3: Context Diagram": ("smartqms-context-diagram.png", "Updated Context Diagram of Smart QMS"),
    "Figure 4: Data Flow Diagram": ("smartqms-data-flow-diagram.png", "Updated Level 1 Data Flow Diagram of Smart QMS"),
    "Figure 5: Entity Relationship Diagram": ("smartqms-erd.png", "Updated Entity Relationship Diagram of Smart QMS"),
    "Figure 6: Use Case Diagram": ("smartqms-use-case-diagram.png", "Updated Use Case Diagram of Smart QMS"),
}

ENTITY_DESCRIPTIONS = [
    "USER - Primary Key: user_id. Stores account details for clients, service staff, and administrators, including role, contact information, client type, and account status. One user can own many queue tickets.",
    "STAFF - Primary Key: staff_id; Foreign Key: user_id. Extends a user account with staff-specific details such as department and shift. A staff member may be assigned to a service window during operations.",
    "HEALTH_SERVICE - Primary Key: service_id. Stores each available health service, its queue mode, estimated duration, and active status. A service can be selected by many queue tickets and supported by one or more windows.",
    "SERVICE_WINDOW - Primary Key: window_id; Foreign Keys: staff_id and service_id. Represents an active physical counter, its current operator, assigned service where applicable, and open, busy, or closed status.",
    "COUNTER_SERVICE - Composite Primary and Foreign Keys: counter_id and service_id. Maps the health services that a service window can support, allowing a window to be configured for more than one service.",
    "QUEUE_TICKET - Primary Key: ticket_id; Foreign Keys: user_id, service_id, and window_id. Records the client's reference number, QR code, ticket number, queue mode, status, check-in information, and lifecycle timestamps.",
    "WAIT_TIME_LOG - Primary Key: log_id; Foreign Key: ticket_id. Stores the queue features, predicted waiting time, actual waiting and service times, confidence, and model version used to evaluate prediction results.",
    "NOTIFICATION - Primary Key: notif_id; Foreign Keys: ticket_id and user_id. Records browser, SMS, or email alerts, including the message, channel, delivery status, read status, and sent time.",
    "FEEDBACK - Primary Key: feedback_id; Foreign Keys: ticket_id, user_id, window_id, and service_id. Stores one completed-ticket rating and optional client comment for service-quality reporting.",
]

USE_CASE_ITEMS = [
    ("Client Use Cases:", True),
    ("Create an account or sign in through the web portal.", False),
    ("Join a queue or schedule an available health service and receive a ticket reference and QR code.", False),
    ("Track the ticket number, queue status, and predicted waiting time.", False),
    ("Receive queue or turn alerts and submit feedback after a completed service.", False),
    ("Service Staff Use Cases:", True),
    ("Check in a client using a QR code, reference number, or walk-in process.", False),
    ("Open or update the assigned service window and call or recall the next ticket.", False),
    ("Start, complete, skip, or void the current ticket according to the queue workflow.", False),
    ("Administrator Use Cases:", True),
    ("Manage client and staff accounts.", False),
    ("Configure health services, service windows, and supported service assignments.", False),
    ("Review queue analytics and generate or export operational reports.", False),
]

PROCEDURE_PARAGRAPHS = [
    "The system is evaluated using an ISO/IEC 25010-based questionnaire. The criteria below are limited to functions that evaluators can directly perform or observe in Smart QMS.",
    "Functional Suitability - Evaluates whether a client can sign in, select a service, join or schedule a queue, receive a ticket reference and QR code, and view the resulting ticket information.",
    "Reliability - Evaluates whether the same queue ticket, status, and service-window result remain consistent after client and staff actions, including preventing more than one active ticket for a client.",
    "Usability - Evaluates whether clients, service staff, and administrators can understand the forms, buttons, messages, queue information, and navigation without unnecessary difficulty.",
    "Performance Efficiency - Evaluates whether ticket issuance, queue updates, service-window actions, and public queue information appear promptly during normal use.",
    "Security - Evaluates whether Smart QMS requires the appropriate sign-in and shows only the pages and actions permitted for the signed-in user role.",
]

EVALUATION_ROWS = {
    1: "1. Functional Suitability",
    2: "1.1 The system allows a client to select an available service and receive one valid queue ticket with a reference number and QR code.",
    3: "1.2 The system shows the client's ticket number, current status, queue information, and estimated waiting time after a ticket is issued.",
    4: "2. Reliability",
    5: "2.1 After a staff action, the ticket status is shown correctly in the staff, client, and public display views.",
    6: "2.2 The system keeps a client from obtaining a second active ticket while an existing ticket is waiting or being served.",
    7: "3. Performance Efficiency",
    8: "3.1 A ticket is issued and displayed promptly after a client submits a valid queue request.",
    9: "3.2 Calling, starting, or completing a ticket updates the queue and service-window status promptly.",
    10: "4. Usability",
    11: "4.1 The system is easy to learn and use for clients, service staff, and administrators without advanced technical knowledge.",
    12: "4.2 The system presents clear buttons, forms, messages, ticket information, and navigation on desktop and mobile screens.",
    13: "5. Security",
    14: "5.1 The system shows only the pages and actions that match the signed-in user's role.",
    15: "5.2 The system requires sign-in before a user can view or change account details and authorized queue-management functions.",
}


def remove_paragraph(paragraph):
    paragraph._element.getparent().remove(paragraph._element)


def set_text(paragraph, text):
    ppr = deepcopy(paragraph._p.pPr) if paragraph._p.pPr is not None else None
    paragraph._p.clear_content()
    if ppr is not None:
        paragraph._p.insert(0, ppr)
    run = paragraph.add_run(text)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.highlight_color = None
    return paragraph


def paragraph_index(document, exact_text):
    return next(index for index, paragraph in enumerate(document.paragraphs) if paragraph.text.strip() == exact_text)


document = Document(SOURCE)

# Replace the four existing inline diagrams with the cleaner, consistent figures.
for caption, (filename, description) in FIGURES.items():
    image_paragraph = document.paragraphs[paragraph_index(document, caption) - 1]
    image_paragraph._p.clear_content()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image = image_paragraph.add_run().add_picture(str(DIAGRAMS / filename), width=Inches(6.5))
    image._inline.docPr.set("title", caption)
    image._inline.docPr.set("descr", description)

# Keep the surrounding diagram explanations current and concise.
set_text(document.paragraphs[paragraph_index(document, "Context Diagram") + 1],
         "The Context Diagram shows the Smart QMS as one system and identifies the Client, Service Staff, Administrator, Public Display, and external prediction and notification services that exchange data with it.")
set_text(document.paragraphs[paragraph_index(document, "Data Flow Diagram") + 1],
         "The DFD presents the system's major internal processes: client and ticket management, queue and window operations, and administration, analytics, and reports. It also shows the operational data store and the external prediction, notification, and public-display flows.")
set_text(document.paragraphs[paragraph_index(document, "Entity Relationship Diagram") + 1],
         "The ERD shows the core data entities used by Smart QMS and the primary relationships that support account management, service routing, queue tickets, waiting-time prediction, notifications, and feedback.")
set_text(document.paragraphs[paragraph_index(document, "Use Case Diagram") + 1],
         "The Use Case Diagram summarizes the functions available to Clients, Service Staff, and Administrators in the current Smart QMS implementation.")

# Replace the old entity descriptions with the current data model descriptions.
entity_heading = paragraph_index(document, "Entity Descriptions:")
use_case_heading = paragraph_index(document, "Use Case Diagram")
existing_entity_paragraphs = list(document.paragraphs[entity_heading + 1:use_case_heading])
entity_style = existing_entity_paragraphs[0].style if existing_entity_paragraphs else document.styles["Normal"]
use_case_paragraph = document.paragraphs[use_case_heading]
for description in ENTITY_DESCRIPTIONS:
    paragraph = use_case_paragraph.insert_paragraph_before(style=entity_style)
    set_text(paragraph, description)
for paragraph in existing_entity_paragraphs:
    remove_paragraph(paragraph)

# Replace the role-specific use-case text with descriptions that match the implemented flows.
use_case_start = paragraph_index(document, "Client Use Cases:")
development_tools = paragraph_index(document, "Development Tools")
existing_use_case_paragraphs = list(document.paragraphs[use_case_start:development_tools])
heading_style = existing_use_case_paragraphs[0].style
body_style = existing_use_case_paragraphs[1].style
development_paragraph = document.paragraphs[development_tools]
for text, is_heading in USE_CASE_ITEMS:
    paragraph = development_paragraph.insert_paragraph_before(style=heading_style if is_heading else body_style)
    set_text(paragraph, text)
for paragraph in existing_use_case_paragraphs:
    remove_paragraph(paragraph)

# Align the evaluation procedure with the five visible, directly testable criteria.
procedure_start = paragraph_index(document, "Evaluation Procedure")
evaluation_tool = paragraph_index(document, "Evaluation Tool")
procedure_targets = document.paragraphs[procedure_start + 1:evaluation_tool]
for paragraph, text in zip(procedure_targets, PROCEDURE_PARAGRAPHS):
    set_text(paragraph, text)
for paragraph in procedure_targets[len(PROCEDURE_PARAGRAPHS):]:
    remove_paragraph(paragraph)

# Make every questionnaire criterion observable in the working application.
evaluation_table = next(table for table in document.tables if table.cell(0, 0).text.strip() == "CRITERIA")
for row_index, text in EVALUATION_ROWS.items():
    set_text(evaluation_table.cell(row_index, 0).paragraphs[0], text)

document.save(OUTPUT)
print(OUTPUT)
