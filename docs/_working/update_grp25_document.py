from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, RGBColor

SOURCE = Path(r"C:\xampp\htdocs\smartqms\docs\_working\Grp25-IT225-Chapters123-widcomments-July9-2026.original.docx")
OUTPUT = Path(r"C:\xampp\htdocs\smartqms\docs\Grp25-IT225-Chapters123-widcomments-July9-2026-updated.docx")
DIAGRAM = Path(r"C:\xampp\htdocs\smartqms\docs\diagrams\smartqms-data-flow-diagram.png")

EVALUATION_PROCEDURE = {
    "Functional Sustainability": (
        "Functional Suitability - The extent to which Smart QMS provides the required functions accurately and completely, "
        "including client registration, service selection, booking or queue joining, ticket and QR-code issuance, live queue "
        "updates, wait-time prediction, notifications, service-window operations, analytics, and reports."
    ),
    "Reliability": (
        "Reliability - The ability of Smart QMS to maintain correct queue, ticket, reservation, and service-window records "
        "while performing consistent queue operations and using safe fallback processing when optional prediction or notification services are unavailable."
    ),
    "Usability": (
        "Usability - The degree to which clients, service staff, and administrators can learn and use the Smart QMS interfaces "
        "efficiently, with clear navigation, forms, status messages, queue information, and responsive layouts."
    ),
    "Performance Efficiency": (
        "Performance Efficiency - The ability of Smart QMS to issue tickets, refresh queue information, process staff actions, "
        "and provide waiting-time estimates and reports within an acceptable response time while using system resources efficiently."
    ),
    "Security": (
        "Security - The capability of Smart QMS to protect personal, account, queue, and operational data through authenticated sessions, "
        "role-based access control, CSRF protection, input validation, secure database access, and controlled access to authorized functions."
    ),
    "Maintainability": (
        "Maintainability - The ease with which authorized developers can understand, test, correct, and enhance the separated PHP, JavaScript, "
        "database, notification, reporting, and machine-learning components without disrupting core queue operations."
    ),
}

EVALUATION_TOOL = {
    "2.1 (what item in Reliability you want evaluated?)": (
        "2.1 The system consistently records and updates queue tickets, check-ins, service windows, and transaction statuses accurately throughout the service process."
    ),
    "2.2 (what item in Reliability you want evaluated?)": (
        "2.2 The system continues to provide valid ticket and queue operations when optional prediction or notification services are unavailable by using built-in fallback processes."
    ),
    "3.1 (what item in Efficiency you want evaluated?)": (
        "3.1 The system processes ticket issuance, real-time queue updates, and service-window actions within an acceptable response time during normal use."
    ),
    "3.2 (what item in Efficiency you want evaluated?)": (
        "3.2 The system displays current queue status and estimated waiting time with minimal delay while using resources efficiently."
    ),
    "6.1  (what item in Maintainability you want evaluated?)": (
        "6.1 The system's modules and services are organized by responsibility, making faults, changes, and enhancements easier to identify and manage."
    ),
    "6.2 (what item in Maintainability you want evaluated?)": (
        "6.2 Authorized maintainers can test, update, or replace individual modules, including reports, notification providers, and the prediction service, without disrupting core queue operations."
    ),
}


def remove_paragraph(paragraph):
    paragraph._element.getparent().remove(paragraph._element)


def set_clean_text(paragraph, text):
    ppr = deepcopy(paragraph._p.pPr) if paragraph._p.pPr is not None else None
    paragraph._p.clear_content()
    if ppr is not None:
        paragraph._p.insert(0, ppr)
    run = paragraph.add_run(text)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.highlight_color = None
    return run


document = Document(SOURCE)

# Replace the explicit DFD placeholder with the verified diagram and remove the old sample figure.
placeholder = next(p for p in document.paragraphs if "paste the illistration" in p.text)
placeholder._p.clear_content()
placeholder.alignment = WD_ALIGN_PARAGRAPH.CENTER
diagram_run = placeholder.add_run()
inline_shape = diagram_run.add_picture(str(DIAGRAM), width=Inches(6.5))
inline_shape._inline.docPr.set("descr", "Level 1 Data Flow Diagram of the Smart Queue Management System")
inline_shape._inline.docPr.set("title", "Smart QMS Data Flow Diagram")

for index, paragraph in list(enumerate(document.paragraphs)):
    if paragraph.text.strip() == "<below is an example of a simple DFD>":
        remove_paragraph(paragraph)
    elif index == 209 and len(paragraph._p.xpath(".//w:drawing")):
        remove_paragraph(paragraph)

# Replace all marked Evaluation Procedure definitions with project-specific ISO/IEC 25010 descriptions.
for paragraph in document.paragraphs:
    normalized = paragraph.text.replace("�", "-").strip()
    for label, replacement in EVALUATION_PROCEDURE.items():
        if normalized.startswith(label):
            set_clean_text(paragraph, replacement)
            break

# Complete the red-placeholder evaluation items in the ISO/IEC 25010 questionnaire.
for table in document.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                current = paragraph.text.strip()
                if current in EVALUATION_TOOL:
                    set_clean_text(paragraph, EVALUATION_TOOL[current])

document.save(OUTPUT)
print(OUTPUT)
