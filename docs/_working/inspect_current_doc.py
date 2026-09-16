from docx import Document

SOURCE = r"C:\xampp\htdocs\smartqms\docs\_working\Grp25-current.docx"
NEEDLES = (
    "Context Diagram", "Data Flow Diagram", "Entity Relationship Diagram", "Entity Descriptions:",
    "Use Case Diagram", "Client Use Cases:", "Evaluation Procedure", "Evaluation Tool", "Maintainability",
)

document = Document(SOURCE)
for index, paragraph in enumerate(document.paragraphs):
    text = paragraph.text.strip()
    if any(needle in text for needle in NEEDLES):
        print(index, repr(text), "drawings=", len(paragraph._p.xpath(".//w:drawing")))

print("tables", len(document.tables), "inline_shapes", len(document.inline_shapes))
for table_index, table in enumerate(document.tables):
    values = [paragraph.text.strip() for row in table.rows for cell in row.cells for paragraph in cell.paragraphs]
    if any(value in {"CRITERIA", "Maintainability", "3. Efficiency", "3. Performance Efficiency"} for value in values):
        print("evaluation_table", table_index)
        for row_index, row in enumerate(table.rows):
            print(row_index, [cell.text.replace("\n", " | ") for cell in row.cells])

for index in range(195, 241):
    paragraph = document.paragraphs[index]
    print("context", index, repr(paragraph.text.strip()), "drawings=", len(paragraph._p.xpath(".//w:drawing")))

for index in range(246, 254):
    print("procedure", index, repr(document.paragraphs[index].text.strip()))
