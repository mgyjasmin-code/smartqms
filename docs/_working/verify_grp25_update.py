from pathlib import Path
from zipfile import ZipFile

from docx import Document

ORIGINAL = Path(r"C:\xampp\htdocs\smartqms\docs\_working\Grp25-IT225-Chapters123-widcomments-July9-2026.original.docx")
UPDATED = Path(r"C:\xampp\htdocs\smartqms\docs\Grp25-IT225-Chapters123-widcomments-July9-2026-updated.docx")

document = Document(UPDATED)
paragraph_text = "\n".join(p.text for p in document.paragraphs)
table_text = "\n".join(
    p.text for table in document.tables for row in table.rows for cell in row.cells for p in cell.paragraphs
)
all_text = paragraph_text + "\n" + table_text

required = [
    "Functional Suitability - The extent to which Smart QMS",
    "The system consistently records and updates queue tickets",
    "The system processes ticket issuance, real-time queue updates",
    "Authorized maintainers can test, update, or replace individual modules",
]
removed = [
    "paste the illistration",
    "below is an example of a simple DFD",
    "what item in Reliability",
    "what item in Efficiency",
    "what item in Maintainability",
]

for item in required:
    print("required", item, item in all_text)
for item in removed:
    print("removed", item, item not in all_text)

dfd_index = next(i for i, p in enumerate(document.paragraphs) if p.text.strip() == "Figure 4: Data Flow Diagram")
prior = document.paragraphs[dfd_index - 1]
print("dfd_picture_before_caption", len(prior._p.xpath(".//w:drawing")) == 1)
print("inline_shapes", len(document.inline_shapes))

for source, name in [(ORIGINAL, "original"), (UPDATED, "updated")]:
    with ZipFile(source) as archive:
        names = archive.namelist()
        comments = archive.read("word/comments.xml") if "word/comments.xml" in names else b""
        print(name, "comments_xml", comments.count(b"<w:comment "), "media", len([n for n in names if n.startswith("word/media/")]))
